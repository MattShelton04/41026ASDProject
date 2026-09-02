"""Strict ABS SEIFA 2021 Suburbs and Localities workbook parser."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

SEIFA_SHEET = "Table 1"
SEIFA_RELEASE_YEAR = 2021
NSW_SAL_PREFIX = "1"
_GROUP_HEADER = (
    None,
    None,
    "Index of Relative Socio-economic Disadvantage",
    None,
    "Index of Relative Socio-economic Advantage and Disadvantage",
    None,
    "Index of Economic Resources",
    None,
    "Index of Education and Occupation",
    None,
    None,
)
_COLUMN_HEADER = (
    "2021 Suburbs and Localities (SAL) Code",
    "2021 Suburbs and Localities (SAL) Name",
    "Score",
    "Decile",
    "Score",
    "Decile",
    "Score",
    "Decile",
    "Score",
    "Decile",
    "Usual Resident Population",
)


class SeifaWorkbookError(ValueError):
    """The publisher workbook does not match the registered 2021 SAL contract."""


@dataclass(frozen=True, slots=True)
class SeifaSalRecord:
    sal_code: str
    sal_name: str
    locality_name: str
    state: str
    reference_year: int
    irsd_score: float | None
    irsd_australia_decile: int | None
    irsad_score: float | None
    irsad_australia_decile: int | None
    ier_score: float | None
    ier_australia_decile: int | None
    ieo_score: float | None
    ieo_australia_decile: int | None
    usual_resident_population: int


def parse_seifa_sal_xlsx(content: bytes, *, state: str = "NSW") -> tuple[SeifaSalRecord, ...]:
    """Parse the complete NSW subset from the official national SAL data cube."""
    if state != "NSW":
        raise SeifaWorkbookError("SEIFA acquisition is registered only for NSW")
    if not content.startswith(b"PK"):
        raise SeifaWorkbookError("SEIFA source is not an XLSX workbook")
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (InvalidFileException, OSError, ValueError) as exc:
        raise SeifaWorkbookError("SEIFA source is not a readable XLSX workbook") from exc
    try:
        if SEIFA_SHEET not in workbook.sheetnames:
            raise SeifaWorkbookError("SEIFA workbook is missing Table 1")
        sheet = workbook[SEIFA_SHEET]
        group_header = tuple(cell.value for cell in sheet[5][:11])
        column_header = tuple(cell.value for cell in sheet[6][:11])
        if group_header != _GROUP_HEADER or column_header != _COLUMN_HEADER:
            raise SeifaWorkbookError("SEIFA Table 1 headers do not match the registered layout")

        records: list[SeifaSalRecord] = []
        seen: set[str] = set()
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=7, max_col=11, values_only=True), start=7
        ):
            raw_code = row[0]
            if raw_code is None:
                continue
            if isinstance(raw_code, bool) or not isinstance(raw_code, (int, float)):
                # The official workbook ends with a copyright footer after all data rows.
                if str(raw_code).lstrip().startswith("©") or "Commonwealth of Australia" in str(
                    raw_code
                ):
                    continue
                raise SeifaWorkbookError(f"SEIFA row {row_number} has an invalid SAL code")
            if int(raw_code) != raw_code:
                raise SeifaWorkbookError(f"SEIFA row {row_number} has an invalid SAL code")
            sal_code = f"{int(raw_code):05d}"
            if not sal_code.startswith(NSW_SAL_PREFIX):
                continue
            if sal_code in seen:
                raise SeifaWorkbookError(f"SEIFA row {row_number} duplicates SAL code {sal_code}")
            seen.add(sal_code)
            sal_name = _required_text(row[1], row_number, "SAL name")
            locality_name = _normalise_nsw_sal_name(sal_name)
            pairs = (
                _index_pair(row[2], row[3], row_number, "IRSD"),
                _index_pair(row[4], row[5], row_number, "IRSAD"),
                _index_pair(row[6], row[7], row_number, "IER"),
                _index_pair(row[8], row[9], row_number, "IEO"),
            )
            population = _required_integer(
                row[10], row_number, "usual resident population", minimum=0
            )
            records.append(
                SeifaSalRecord(
                    sal_code=sal_code,
                    sal_name=sal_name,
                    locality_name=locality_name,
                    state="NSW",
                    reference_year=SEIFA_RELEASE_YEAR,
                    irsd_score=pairs[0][0],
                    irsd_australia_decile=pairs[0][1],
                    irsad_score=pairs[1][0],
                    irsad_australia_decile=pairs[1][1],
                    ier_score=pairs[2][0],
                    ier_australia_decile=pairs[2][1],
                    ieo_score=pairs[3][0],
                    ieo_australia_decile=pairs[3][1],
                    usual_resident_population=population,
                )
            )
        if not records:
            raise SeifaWorkbookError("SEIFA workbook contains no NSW SAL records")
        return tuple(sorted(records, key=lambda item: item.sal_code))
    finally:
        workbook.close()


def _normalise_nsw_sal_name(value: str) -> str:
    suffix = " (NSW)"
    name = value[: -len(suffix)] if value.upper().endswith(suffix) else value
    return " ".join(name.upper().split())


def _required_text(value: Any, row: int, field: str) -> str:
    text = str(value).strip() if value is not None else ""
    if not text:
        raise SeifaWorkbookError(f"SEIFA row {row} {field} is required")
    return text


def _index_pair(score: Any, decile: Any, row: int, label: str) -> tuple[float | None, int | None]:
    missing_score = score in {None, "-"}
    missing_decile = decile in {None, "-"}
    if missing_score or missing_decile:
        if missing_score and missing_decile:
            return None, None
        raise SeifaWorkbookError(f"SEIFA row {row} {label} score and decile must both be present")
    if isinstance(score, bool) or not isinstance(score, (int, float)):
        raise SeifaWorkbookError(f"SEIFA row {row} {label} score must be numeric")
    parsed_decile = _required_integer(decile, row, f"{label} decile", minimum=1)
    if parsed_decile > 10:
        raise SeifaWorkbookError(f"SEIFA row {row} {label} decile must be between 1 and 10")
    return float(score), parsed_decile


def _required_integer(value: Any, row: int, field: str, *, minimum: int) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)) or int(value) != value:
        raise SeifaWorkbookError(f"SEIFA row {row} {field} must be an integer")
    parsed = int(value)
    if parsed < minimum:
        raise SeifaWorkbookError(f"SEIFA row {row} {field} is outside the registered range")
    return parsed
